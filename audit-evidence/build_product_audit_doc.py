from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(r"E:\codex\libtv-studio-complete-v1.3")
EVIDENCE_DIR = ROOT / "audit-evidence" / "2026-08-20"
OUT = ROOT / "audit-evidence" / "LibTV-Studio-Product-Audit-2026-08-20.docx"

BLUE = "2E74B5"
NAVY = "0B2545"
INK = "1F2937"
MUTED = "667085"
LIGHT = "F2F4F7"
CALLOUT = "F4F6F9"
P1_RED = "9B1C1C"
P1_FILL = "FCE8E6"
P2_GOLD = "7A5A00"
P2_FILL = "FFF8E8"
P3_GRAY = "475467"
P3_FILL = "F2F4F7"
GREEN = "166534"
GREEN_FILL = "ECFDF3"


def set_run_font(run, name="Calibri", size=11, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D0D5DD", size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = True
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        st.font.size = Pt(11)
        st.paragraph_format.left_indent = Inches(0.5)
        st.paragraph_format.first_line_indent = Inches(-0.25)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.line_spacing = 1.167

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("LibTV Studio  |  Product Audit")
    set_run_font(run, size=9, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.paragraph_format.space_before = Pt(0)
    run = footer.add_run("Internal working document  •  2026-08-20")
    set_run_font(run, size=9, color=MUTED)


def add_title_block(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("PRODUCT AUDIT")
    set_run_font(r, size=10, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("LibTV Studio 产品审计")
    set_run_font(r, size=28, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("真实用户完整走查 | 资深 PM + UX + QA 视角 | 基于实际操作证据")
    set_run_font(r, size=13, color=MUTED)

    metadata = [
        ("审计日期", "2026-08-20"),
        ("产品范围", "LibTV Studio 本地桌面/浏览器创作工作流"),
        ("走查链路", "新建画布 → 文本 → 分镜 → 图片/视频 → 资产/提示词 → Agent Skill → 时间线 → 导出"),
        ("变更边界", "本次仅审计与记录，未修改产品代码"),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    set_table_geometry(table, [2100, 7260])
    set_table_borders(table, color="D0D5DD", size="4")
    for row, (label, value) in zip(table.rows, metadata):
        set_cell_shading(row.cells[0], LIGHT)
        p0 = row.cells[0].paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        set_run_font(r0, size=10, color=NAVY, bold=True)
        p1 = row.cells[1].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(value)
        set_run_font(r1, size=10, color=INK)


def add_callout(doc, label, text, fill=CALLOUT, label_color=NAVY):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    set_table_borders(table, color="D0D5DD", size="4")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + "  ")
    set_run_font(r, size=11, color=label_color, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_labeled_para(doc, label, text, label_color=NAVY):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(label + "：")
    set_run_font(r, size=11, color=label_color, bold=True)
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    set_run_font(r, size=11, color=INK)
    return p


def add_finding(doc, code, title, evidence, impact, recommendation):
    p = doc.add_paragraph(style="Heading 3")
    r = p.add_run(f"{code}  {title}")
    set_run_font(r, size=12, color=P1_RED, bold=True)
    add_labeled_para(doc, "实际证据", evidence, label_color=P1_RED)
    add_labeled_para(doc, "用户/业务影响", impact, label_color=P1_RED)
    add_labeled_para(doc, "建议", recommendation, label_color=GREEN)


def add_summary(doc):
    doc.add_heading("一、执行摘要", level=1)
    add_callout(
        doc,
        "总体判断",
        "核心创作链路可以被首次用户理解并走通，但生成任务的状态、依赖失败恢复和空状态边界还不够可信。产品现在最需要的不是重做交互范式，而是建立一套跨文本、图片、视频和 Skill 的统一任务状态真相源。",
        fill=CALLOUT,
    )

    table = doc.add_table(rows=2, cols=5)
    set_table_geometry(table, [1500, 1800, 1800, 1800, 2460])
    set_table_borders(table, color="D0D5DD", size="4")
    headers = [("等级", "数量", "含义", "当前判断", "处理节奏")]
    for row in table.rows[:1]:
        set_repeat_table_header(row)
        for i, text in enumerate(headers[0]):
            set_cell_shading(row.cells[i], LIGHT)
            p = row.cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_run_font(r, size=10, color=NAVY, bold=True)
    values = [("P0", "0", "阻断/数据损失", "未发现", "无需作为发布阻塞")]
    # Add P0 as the first row of data and use separate compact rows below.
    for i, text in enumerate(values[0]):
        p = table.rows[1].cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=10, color=P1_RED if i == 0 else INK, bold=(i == 0))

    for level, count, meaning, judgment, cadence, fill, color in [
        ("P1", "5", "核心链路风险", "生成状态、依赖恢复、导出边界、Skill 状态", "本版本优先", P1_FILL, P1_RED),
        ("P2", "4", "明显可用性问题", "首屏缩放、横向滚动、空态 CTA、Skill 路径", "紧随其后", P2_FILL, P2_GOLD),
        ("P3", "3", "表达/效率优化", "导航、模型反馈、建议入口", "排期优化", P3_FILL, P3_GRAY),
    ]:
        row = table.add_row()
        for i, text in enumerate((level, count, meaning, judgment, cadence)):
            p = row.cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_run_font(r, size=10, color=color if i == 0 else INK, bold=(i == 0))
            if i == 0:
                set_cell_shading(row.cells[i], fill)
    set_table_geometry(table, [1500, 1800, 1800, 1800, 2460])

    doc.add_heading("正向观察", level=2)
    for text in [
        "空白画布能快速开始，首次用户不需要先理解复杂项目结构。",
        "文本 → 图片/视频的节点关系直观，分镜节点是产品的强概念。",
        "图片依赖可以自动触发；取消后提供“重新生成”，恢复路径方向正确。",
        "资产库、提示词库与快捷键帮助具备复用价值，能降低后续学习成本。",
        "Agent 只读问答与 Skill 写入的边界清晰，产品方向是对的。",
    ]:
        add_bullet(doc, text)


def add_method(doc):
    doc.add_heading("二、走查方法与流程证据", level=1)
    add_labeled_para(doc, "用户视角", "按第一次使用产品的真实用户路径操作，不预先阅读源码、不假设内部状态正确。")
    add_labeled_para(doc, "操作覆盖", "创建/打开项目、创建空白画布、添加文本节点、填写提示词、等待与取消生成、创建脚本与分镜、触发图片/视频依赖、查看资产/提示词、运行 Agent Skill、查看时间线、尝试导出。")
    add_labeled_para(doc, "异常覆盖", "等待超过 1 分钟、取消、上游失败、下游停止、空时间线导出、重复/陈旧校验、面板切换与空态。")
    add_labeled_para(doc, "证据规则", "每条问题均来自实际操作时看到的界面或状态，截图原文件保存在 audit-evidence/2026-08-20。")

    doc.add_heading("流程结论", level=2)
    steps = [
        ("01", "打开/新建", "顺畅：可以快速进入现有项目或新建空白画布。"),
        ("02", "首次放置文本", "可用但发现性弱：默认缩放约 25%，节点和下一步不够突出。"),
        ("03", "文本生成", "风险：观察到生成状态与结果/可操作性表现不一致。"),
        ("04", "脚本/分镜", "基本可用：全屏表格能完成操作，但横向内容超出视口。"),
        ("05", "图片生成", "风险：持续超过 1 分钟，无明确超时与失败反馈。"),
        ("06", "视频依赖", "逻辑方向正确：上游失败会阻止下游，但恢复动作不足。"),
        ("07", "资产/提示词", "可用：空态与复用入口相对清楚。"),
        ("08", "Agent Skill", "部分可用：Skill 已排队/创建节点后仍残留“缺少必填输入”。"),
        ("09", "时间线", "偏弱：空态缺少“添加片段/从分镜导入”的明确 CTA。"),
        ("10", "导出", "边界不清：空时间线仍可触发导出，随后出现技术化英文错误。"),
    ]
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [900, 2100, 6360])
    set_table_borders(table, color="D0D5DD", size="4")
    set_repeat_table_header(table.rows[0])
    for i, text in enumerate(("步骤", "流程节点", "实际观察")):
        set_cell_shading(table.rows[0].cells[i], LIGHT)
        p = table.rows[0].cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=10, color=NAVY, bold=True)
    for step, name, result in steps:
        row = table.add_row()
        for i, text in enumerate((step, name, result)):
            p = row.cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_run_font(r, size=10, color=INK, bold=(i == 0))
    set_table_geometry(table, [900, 2100, 6360])


def add_findings(doc):
    doc.add_page_break()
    doc.add_heading("三、P1 问题：核心链路优先修复", level=1)
    add_callout(doc, "P1 共性根因", "任务系统、依赖系统、前端反馈没有共享同一份用户可见的状态真相。先统一状态机，再处理视觉细节，收益会覆盖文本、图片、视频与 Skill 四条链路。", fill=P1_FILL, label_color=P1_RED)
    doc.add_heading("P1 问题明细", level=2)

    add_finding(doc, "P1-01", "生成状态机不一致", "在文本生成期间观察到“生成中”与仍可继续操作的矛盾表现；等待后结果与状态仍不一致。对应截图：07-text-generating、08-text-generating-wait、09-text-result-still-generating。", "用户无法判断任务是否已成功、是否应该再次提交，容易造成重复操作和对结果可靠性的怀疑。", "统一 queued / running / success / failed / cancelled 五态；每个状态只保留一组主动作，并让节点、面板、通知使用同一状态源。")
    add_finding(doc, "P1-02", "图片生成无超时与明确失败", "图片节点持续超过 1 分钟，界面没有明确超时、预计等待阶段或失败原因；直到取消后才回到可操作状态。对应截图：13-image-generating、14-image-cancelled。", "用户会把“等待”误认为“卡死”，无法决定继续等待、调整参数还是换模型。", "增加超时阈值、阶段/进度反馈、保留原参数的“重试”、可理解的失败原因与取消后的下一步。")
    add_finding(doc, "P1-03", "上游失败后的依赖恢复弱", "图片失败后视频自动停止；不同观察时点出现“已失败”和“仍生成中”的冲突状态。对应截图：16-video-autoruns-image-dependency、17c-upstream-failure-immediate、17-upstream-failure-stops-video。", "下游被动失效，用户不知道先修哪个节点，也没有“修复上游并恢复下游”的一键路径。", "让错误沿依赖图传播：下游显示“等待上游修复”；提供重试上游、恢复依赖链、重新运行下游的明确动作。")
    add_finding(doc, "P1-04", "空时间线仍允许导出", "时间线为空时仍可点击导出；随后出现技术化英文错误，没有在导出入口提前阻止。对应截图：23-timeline-empty、24b-export-error-immediate。", "用户多走一步才知道不能完成目标，错误文案也没有告诉用户下一步应添加什么。", "空时间线时禁用导出按钮，并解释“至少添加一个可导出片段”；误触时使用中文、可行动的提示。")
    add_finding(doc, "P1-05", "Agent Skill 存在陈旧必填校验", "Skill 已排队或已创建节点后，界面仍显示“缺少必填输入”；状态与结果不一致。对应截图：20-skill-required-input、21-agent-skill-queued、22-skill-nodes-on-canvas。", "用户会认为操作没有生效，重复点击或放弃，尤其影响对自动化能力的信任。", "提交成功后立即清掉旧校验；展示 queued / running / success / failed，并提供结果入口与刷新后的一致性。")

    doc.add_heading("QA 验收标准", level=2)
    for text in [
        "任一生成任务在五种状态下都有唯一文案、唯一主按钮、唯一结果判断。",
        "上游失败后，下游不会显示“生成中”；恢复上游后可以明确重试下游。",
        "空时间线导出按钮不可用，辅助文案说明至少需要一个可导出片段。",
        "Skill 提交成功后，旧的必填错误立即消失；刷新后状态仍一致。",
        "断网、超时、取消、重复点击均有可恢复路径，且不会产生重复节点或重复任务。",
    ]:
        add_number(doc, text)


def add_p2_p3(doc):
    doc.add_heading("四、P2 / P3 问题：可用性与表达", level=1)
    doc.add_heading("P2（4 项）", level=2)
    for text in [
        "首次打开画布默认约 25% 缩放，节点和下一步不易发现。建议首次进入自动适配视口，并突出“添加文本节点 → 生成”。",
        "脚本全屏表格可用，但横向内容远超视口，滚动提示弱。建议提升列收缩/换行策略，并在首次出现时提示可横向滚动。",
        "时间线空态缺少“添加片段 / 从分镜导入”的明确 CTA。建议把空态变成下一步选择，而不是仅展示说明。",
        "Skill“使用”路径直接跳 Agent，缺少执行前说明与预期结果。建议在跳转前说明输入、生成结果和预计状态。",
    ]:
        add_bullet(doc, text)

    doc.add_heading("P3（3 项）", level=2)
    for text in [
        "中等宽度下顶部导航出现图标化，缺少文字或 Tooltip 辅助。",
        "模型/API 面板缺少“当前生效模型”和“测试连接”反馈。",
        "“尝试”建议的点击区域与点击目的不够清晰。",
    ]:
        add_bullet(doc, text)


def add_pm_five(doc):
    doc.add_page_break()
    doc.add_heading("五、如果下个版本只允许改 5 件事", level=1)
    add_callout(doc, "PM 排序原则", "优先提升核心任务完成率、用户信任与失败恢复，而不是先做低风险的视觉润色。", fill=GREEN_FILL, label_color=GREEN)
    five = [
        ("01", "统一生成状态机", "覆盖文本、图片、视频、Skill；让状态、主按钮、结果入口一致。"),
        ("02", "依赖失败与一键恢复", "明确上游/下游关系；失败后给出修复顺序、重试与继续运行。"),
        ("03", "时间线空态与导出边界", "空时间线提前禁用导出；下一步变成添加片段或从分镜导入。"),
        ("04", "首次使用画布引导", "首次进入自动适配缩放，并突出添加文本节点 → 生成的第一步。"),
        ("05", "Skill 执行反馈与校验修复", "提交成功即清掉旧错误；展示排队、运行、完成/失败与结果。"),
    ]
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [900, 2700, 5760])
    set_table_borders(table, color="B8C7E6", size="6")
    set_repeat_table_header(table.rows[0])
    for i, text in enumerate(("优先级", "改动", "为什么现在做")):
        set_cell_shading(table.rows[0].cells[i], "E8EEF5")
        p = table.rows[0].cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=10, color=NAVY, bold=True)
    for order, name, why in five:
        row = table.add_row()
        set_cell_shading(row.cells[0], "E8EEF5")
        for i, text in enumerate((order, name, why)):
            p = row.cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_run_font(r, size=10, color=NAVY if i < 2 else INK, bold=(i < 2))
    set_table_geometry(table, [900, 2700, 5760])

    doc.add_heading("建议版本目标", level=2)
    for text in [
        "用户能在任意时刻回答：现在发生了什么、还要等多久、失败后怎么恢复。",
        "用户不会在空时间线或缺少输入时走到不可完成的下一步。",
        "生成失败不会破坏已创建的上游内容，也不会留下互相矛盾的状态。",
    ]:
        add_bullet(doc, text)


def add_evidence(doc):
    doc.add_page_break()
    doc.add_heading("六、截图证据索引", level=1)
    add_labeled_para(doc, "截图位置", "E:\\codex\\libtv-studio-complete-v1.3\\audit-evidence\\2026-08-20")
    add_labeled_para(doc, "说明", "下表为最能支撑主要结论的代表性截图；完整目录在表后列出。截图均来自本次真实操作。")

    evidence = [
        ("03-empty-canvas.png", "新建/空白画布", "支持正向结论：空白画布能快速开始；也可用于观察首次缩放与下一步发现性。"),
        ("08-text-generating-wait.png", "文本生成状态", "支撑 P1-01：等待期间状态与可操作性/结果反馈不够一致。"),
        ("13-image-generating.png", "图片生成等待", "支撑 P1-02：图片任务长时间没有明确超时或失败反馈。"),
        ("17c-upstream-failure-immediate.png", "上游失败/下游停止", "支撑 P1-03：依赖失败时恢复顺序和动作不足。"),
        ("21-agent-skill-queued.png", "Agent Skill 排队", "支撑 P1-05：排队状态与残留的必填校验同时存在。"),
        ("24b-export-error-immediate.png", "空时间线导出错误", "支撑 P1-04：空时间线仍允许触发导出，错误是技术化英文。"),
    ]
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [2650, 1850, 4860])
    set_table_borders(table, color="D0D5DD", size="4")
    set_repeat_table_header(table.rows[0])
    for i, text in enumerate(("截图", "场景", "审计用途")):
        set_cell_shading(table.rows[0].cells[i], LIGHT)
        p = table.rows[0].cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=10, color=NAVY, bold=True)
    for filename, scene, use in evidence:
        row = table.add_row()
        p = row.cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        image_path = EVIDENCE_DIR / filename
        if image_path.exists():
            inline = p.add_run().add_picture(str(image_path), width=Inches(1.55))
            inline._inline.docPr.set("descr", f"审计证据截图：{scene}")
            inline._inline.docPr.set("title", f"Audit evidence: {scene}")
        p2 = row.cells[0].add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        r = p2.add_run(filename)
        set_run_font(r, size=8.5, color=MUTED)
        for idx, text in ((1, scene), (2, use)):
            p = row.cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            set_run_font(r, size=10, color=INK)
    set_table_geometry(table, [2650, 1850, 4860])

    doc.add_heading("完整截图目录（按操作流程）", level=2)
    groups = [
        ("创建与画布", ["01-existing-project.png", "02-new-canvas-modal.png", "03-empty-canvas.png", "04-text-node-collapsed.png", "05-text-node-fit.png", "06-text-input-filled.png"]),
        ("文本与生成", ["07-text-generating.png", "08-text-generating-wait.png", "09-text-result-still-generating.png"]),
        ("脚本与分镜", ["10-script-full-view.png", "11-storyboard-nodes-created.png", "12-storyboard-fit.png"]),
        ("图片与视频依赖", ["13-image-generating.png", "14-image-cancelled.png", "15-text-completed.png", "16-video-autoruns-image-dependency.png", "17-upstream-failure-stops-video.png", "17c-upstream-failure-immediate.png"]),
        ("模型、资产与 Skill", ["18-model-api-panel.png", "19-assets-empty.png", "20-skill-required-input.png", "21-agent-skill-queued.png", "22-skill-nodes-on-canvas.png"]),
        ("时间线、导出与基础交互", ["23-timeline-empty.png", "24-export-empty-error.png", "24b-export-error-immediate.png", "25-add-node-menu.png", "26-shortcuts-help.png"]),
    ]
    for group, files in groups:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(group + "：")
        set_run_font(r, size=10, color=NAVY, bold=True)
        r = p.add_run("、".join(files))
        set_run_font(r, size=9, color=MUTED)

    doc.add_heading("审计边界", level=2)
    add_labeled_para(doc, "未发现", "本次没有观察到 P0 级别的阻断、数据丢失或不可逆破坏问题。")
    add_labeled_para(doc, "未做事项", "没有修改产品源码、没有改变产品配置、没有替用户提交外部服务或发布内容。")
    add_labeled_para(doc, "后续建议", "修复 P1 后，用同一组截图步骤做回归；重点验证任务状态、失败恢复、空态导出和 Skill 状态刷新。")


def main():
    doc = Document()
    style_document(doc)
    add_title_block(doc)
    add_summary(doc)
    add_method(doc)
    add_findings(doc)
    add_p2_p3(doc)
    add_pm_five(doc)
    add_evidence(doc)
    doc.core_properties.title = "LibTV Studio Product Audit"
    doc.core_properties.subject = "Product audit based on real user operation evidence"
    doc.core_properties.author = "Codex"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
